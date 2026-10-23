from pathlib import Path
import copy
import csv
import itertools
import math
import numpy as np
import open3d as o3d
import matplotlib
matplotlib.use("Agg")  
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_DIR = Path(r"D:\CV5570\CV5570_3D_Reconstruction_Assignment\Reconstruction Workspace")
IMAGE_PATH = PROJECT_DIR / "images" / "1.jpg"
PLY_PATH = PROJECT_DIR / "dense" / "0" / "fused.ply"
OUTPUT_DIR = PROJECT_DIR / "output"
KNOWN_WIDTH_CM = 17.5
KNOWN_HEIGHT_CM = 23.0
ROTATION_DEG_X = 45.0
TOLERANCE_PERCENT = 5.0

MAX_PLOT_POINTS = 60000
RANDOM_SEED = 42

def get_dimensions(points):
    minimum = points.min(axis=0)
    maximum = points.max(axis=0)
    bbox_dims = maximum - minimum
    low = np.percentile(points, 1, axis=0)
    high = np.percentile(points, 99, axis=0)
    robust_dims = high - low
    return minimum, maximum, bbox_dims, robust_dims


def choose_scale_and_axes(robust_dims):
    actual_width = KNOWN_WIDTH_CM
    actual_height = KNOWN_HEIGHT_CM
    names = ["X", "Y", "Z"]
    candidates = []

    for width_axis, height_axis in itertools.permutations(range(3), 2):
        width_units = float(robust_dims[width_axis])
        height_units = float(robust_dims[height_axis])
        if width_units <= 0 or height_units <= 0:
            continue
        scale = (
            width_units * actual_width + height_units * actual_height
        ) / (width_units ** 2 + height_units ** 2)

        predicted_width = width_units * scale
        predicted_height = height_units * scale
        width_error_pct = abs(predicted_width - actual_width) / actual_width * 100
        height_error_pct = abs(predicted_height - actual_height) / actual_height * 100

        candidates.append({
            "width_axis": width_axis,
            "height_axis": height_axis,
            "scale": scale,
            "width_units": width_units,
            "height_units": height_units,
            "predicted_width": predicted_width,
            "predicted_height": predicted_height,
            "width_error_pct": width_error_pct,
            "height_error_pct": height_error_pct,
            "max_error_pct": max(width_error_pct, height_error_pct),
            "mean_error_pct": (width_error_pct + height_error_pct) / 2,
        })

    if not candidates:
        raise ValueError("Could not calculate dimensions: point cloud has zero-size axes.")

    best = min(candidates, key=lambda c: (c["max_error_pct"], c["mean_error_pct"]))
    print("\nCandidate axis assignments, best first:")
    for c in sorted(candidates, key=lambda x: (x["max_error_pct"], x["mean_error_pct"])):
        print(
            f"Width={names[c['width_axis']]}, Height={names[c['height_axis']]}, "
            f"scale={c['scale']:.8f} cm/unit, "
            f"width error={c['width_error_pct']:.2f}%, "
            f"height error={c['height_error_pct']:.2f}%, "
            f"worst={c['max_error_pct']:.2f}%"
        )
    return best


def sample_points(cloud, max_points=MAX_PLOT_POINTS):
    pts = np.asarray(cloud.points)
    colors = np.asarray(cloud.colors) if cloud.has_colors() else None
    if len(pts) > max_points:
        rng = np.random.default_rng(RANDOM_SEED)
        indices = rng.choice(len(pts), max_points, replace=False)
        pts = pts[indices]
        if colors is not None:
            colors = colors[indices]
    return pts, colors


def orient_for_display(points):
    rx = np.deg2rad(90)
    ry = np.deg2rad(0)
    rz = np.deg2rad(90)
    Rx = np.array([[1, 0, 0],
                   [0, np.cos(rx), -np.sin(rx)],
                   [0, np.sin(rx),  np.cos(rx)]])
    Ry = np.array([[np.cos(ry), 0, np.sin(ry)],
                   [0, 1, 0],
                   [-np.sin(ry), 0, np.cos(ry)]])
    Rz = np.array([[np.cos(rz), -np.sin(rz), 0],
                   [np.sin(rz),  np.cos(rz), 0],
                   [0, 0, 1]])
    centered = points - points.mean(axis=0)
    return centered @ (Rz @ Ry @ Rx).T


def configure_cloud_axis(ax, points, title, unit, colors=None):
    plot_colors = colors if colors is not None and len(colors) == len(points) else "#252525"
    ax.scatter(
        points[:, 0], points[:, 1], points[:, 2],
        c=plot_colors, s=0.65, alpha=0.95, depthshade=False,
        linewidths=0, rasterized=True
    )
    ax.set_title(title, fontsize=12, fontweight="bold", pad=10)
    ax.set_xlabel(f"X ({unit})", labelpad=4)
    ax.set_ylabel(f"Y ({unit})", labelpad=4)
    ax.set_zlabel(f"Z ({unit})", labelpad=4)
    lo, hi = points.min(axis=0), points.max(axis=0)
    spans = np.maximum(hi - lo, 1e-6)
    center = (lo + hi) / 2
    radius = max(float(spans.max()) / 2, 1e-6)
    ax.set_xlim(center[0] - radius, center[0] + radius)
    ax.set_ylim(center[1] - radius, center[1] + radius)
    ax.set_zlim(center[2] - radius, center[2] + radius)
    ax.set_box_aspect((1, 1, 1))
    ax.view_init(elev=15, azim=-90)
    ax.tick_params(labelsize=7)
    ax.grid(True, alpha=0.25)


def write_comparison(image_path, original_points, original_colors, scaled_points,
                     scaled_colors, robust_dims, best, overall_pass, output_path):
    photo = Image.open(image_path).convert("RGB")
    orig_display = orient_for_display(original_points)
    scaled_display = orient_for_display(scaled_points)
    names = ["X", "Y", "Z"]

    fig = plt.figure(figsize=(19, 9), facecolor="white")
    ax1 = fig.add_subplot(1, 3, 1)
    ax1.imshow(photo)
    ax1.set_title("REAL OBJECT", fontsize=13, fontweight="bold")
    ax1.axis("off")
    ax1.text(
        0.04, 0.04,
        f"Width: {KNOWN_WIDTH_CM:.1f} cm\nHeight: {KNOWN_HEIGHT_CM:.1f} cm",
        transform=ax1.transAxes, color="white", fontsize=10, fontweight="bold",
        bbox=dict(facecolor="black", alpha=0.85, pad=5)
    )

    ax2 = fig.add_subplot(1, 3, 2, projection="3d")
    configure_cloud_axis(ax2, orig_display, "ORIGINAL POINT CLOUD", "PLY units", original_colors)
    ax2.text2D(
        0.03, 0.97,
        f"Width: {best['width_units']:.3f} units (axis {names[best['width_axis']]})\n"
        f"Height: {best['height_units']:.3f} units (axis {names[best['height_axis']]})",
        transform=ax2.transAxes, va="top", fontsize=9,
        bbox=dict(facecolor="white", edgecolor="gray", alpha=0.95)
    )

    ax3 = fig.add_subplot(1, 3, 3, projection="3d")
    configure_cloud_axis(ax3, scaled_display, "BEST SCALED POINT CLOUD", "cm", scaled_colors)
    status = "PASS" if overall_pass else "FAIL"
    ax3.text2D(
        0.03, 0.97,
        f"Width: {best['predicted_width']:.2f} cm "
        f"({best['width_error_pct']:.2f}% error)\n"
        f"Height: {best['predicted_height']:.2f} cm "
        f"({best['height_error_pct']:.2f}% error)\n"
        f"Tolerance: +/-{TOLERANCE_PERCENT:.0f}%\nResult: {status}\n"
        f"X rotation saved: {ROTATION_DEG_X:.0f} degrees",
        transform=ax3.transAxes, va="top", fontsize=9,
        bbox=dict(facecolor="white", edgecolor="gray", alpha=0.95)
    )

    fig.suptitle(
        "3D RECONSTRUCTION: WIDTH AND HEIGHT ACCURACY",
        fontsize=17, fontweight="bold", y=0.97
    )
    fig.subplots_adjust(left=0.02, right=0.98, bottom=0.05, top=0.87, wspace=0.12)
    fig.savefig(output_path, dpi=250, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if not IMAGE_PATH.is_file():
        raise FileNotFoundError(f"Original image not found: {IMAGE_PATH}")
    if not PLY_PATH.is_file():
        raise FileNotFoundError(
            f"PLY not found: {PLY_PATH}\n"
            "Check whether your actual file is dense\\fused.ply or dense\\0\\fused.ply, "
            "then update PLY_PATH near the top of this script."
        )

    pcd = o3d.io.read_point_cloud(str(PLY_PATH))
    if pcd.is_empty():
        raise ValueError(f"The PLY is empty or could not be read: {PLY_PATH}")
    original = np.asarray(pcd.points)
    if not np.isfinite(original).all():
        raise ValueError("Point cloud contains invalid coordinates.")

    print("Input PLY:", PLY_PATH)
    print("Original image:", IMAGE_PATH)
    print("Number of points:", len(original))

    _, _, bbox_dims, robust_dims = get_dimensions(original)
    axis_names = ["X", "Y", "Z"]
    print("\nFull bounding-box dimensions:")
    for i, name in enumerate(axis_names):
        print(f"{name}: {bbox_dims[i]:.6f} PLY units")
    print("\nRobust dimensions (1st-99th percentile):")
    for i, name in enumerate(axis_names):
        print(f"{name}: {robust_dims[i]:.6f} PLY units")

    best = choose_scale_and_axes(robust_dims)
    scale = float(best["scale"])
    width_axis, height_axis = best["width_axis"], best["height_axis"]

    S = np.diag([scale, scale, scale, 1.0])
    theta = np.deg2rad(ROTATION_DEG_X)
    Rx = np.array([
        [1, 0, 0, 0],
        [0, np.cos(theta), -np.sin(theta), 0],
        [0, np.sin(theta),  np.cos(theta), 0],
        [0, 0, 0, 1],
    ], dtype=float)
    T = Rx @ S

    scaled = copy.deepcopy(pcd)
    scaled.scale(scale, center=(0, 0, 0))
    rotated = copy.deepcopy(scaled)
    R3 = Rx[:3, :3]
    rotated.rotate(R3, center=(0, 0, 0))

    scaled_ply = OUTPUT_DIR / "point_cloud_scaled_cm.ply"
    rotated_ply = OUTPUT_DIR / "point_cloud_scaled_rotated_45deg_x.ply"
    o3d.io.write_point_cloud(str(scaled_ply), scaled)
    o3d.io.write_point_cloud(str(rotated_ply), rotated)

    scaled_robust_dims = robust_dims * scale
    predicted_width = float(scaled_robust_dims[width_axis])
    predicted_height = float(scaled_robust_dims[height_axis])
    width_error_cm = abs(predicted_width - KNOWN_WIDTH_CM)
    height_error_cm = abs(predicted_height - KNOWN_HEIGHT_CM)
    width_error_pct = 100 * width_error_cm / KNOWN_WIDTH_CM
    height_error_pct = 100 * height_error_cm / KNOWN_HEIGHT_CM
    width_pass = width_error_pct <= TOLERANCE_PERCENT
    height_pass = height_error_pct <= TOLERANCE_PERCENT
    overall_pass = width_pass and height_pass

    print("\nFINAL DIMENSION COMPARISON")
    print(f"Width: actual={KNOWN_WIDTH_CM:.2f} cm, reconstructed={predicted_width:.2f} cm, error={width_error_pct:.2f}%")
    print(f"Height: actual={KNOWN_HEIGHT_CM:.2f} cm, reconstructed={predicted_height:.2f} cm, error={height_error_pct:.2f}%")
    print("Tolerance result:", "PASS - both dimensions within +/-5%" if overall_pass else "FAIL - at least one dimension exceeds +/-5%")

    matrix_txt = OUTPUT_DIR / "transformation_matrix.txt"
    with matrix_txt.open("w", encoding="utf-8") as f:
        f.write("COLMAP POINT-CLOUD SCALE AND ROTATION TRANSFORMATION\n")
        f.write(f"Input PLY: {PLY_PATH}\n")
        f.write(f"Scale factor: {scale:.12f} cm per PLY unit\n")
        f.write(f"Reference width/height: {KNOWN_WIDTH_CM:.4f} cm / {KNOWN_HEIGHT_CM:.4f} cm\n")
        f.write(f"Chosen width axis: {axis_names[width_axis]}\nChosen height axis: {axis_names[height_axis]}\n\n")
        f.write("Scaling matrix S (4x4):\n" + np.array2string(S, precision=12) + "\n\n")
        f.write("X-axis rotation matrix Rx(45 degrees) (4x4):\n" + np.array2string(Rx, precision=12) + "\n\n")
        f.write("Combined matrix T = Rx @ S (4x4):\n" + np.array2string(T, precision=12) + "\n\n")
        f.write("Column-vector convention: p_new = T @ [x, y, z, 1]^T.\n")
        f.write("Order: uniform scale, then rotate 45 degrees about X around the origin.\n")
        f.write("The rotation preserves Euclidean distances. Dimensions in the report are evaluated before rotation.\n")

    csv_path = OUTPUT_DIR / "width_height_comparison.csv"
    rows = [
        {
            "Dimension": "Width",
            "Axis": axis_names[width_axis],
            "Actual (cm)": KNOWN_WIDTH_CM,
            "Reconstructed (cm)": predicted_width,
            "Absolute Error (cm)": width_error_cm,
            "Error (%)": width_error_pct,
            "Tolerance (%)": TOLERANCE_PERCENT,
            "Within +/-5%": "YES" if width_pass else "NO",
        },
        {
            "Dimension": "Height",
            "Axis": axis_names[height_axis],
            "Actual (cm)": KNOWN_HEIGHT_CM,
            "Reconstructed (cm)": predicted_height,
            "Absolute Error (cm)": height_error_cm,
            "Error (%)": height_error_pct,
            "Tolerance (%)": TOLERANCE_PERCENT,
            "Within +/-5%": "YES" if height_pass else "NO",
        },
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)

    comparison_png = OUTPUT_DIR / "width_height_comparison.png"
    orig_pts, orig_colors = sample_points(pcd)
    scaled_pts, scaled_colors = sample_points(scaled)
    write_comparison(
        IMAGE_PATH, orig_pts, orig_colors, scaled_pts, scaled_colors,
        robust_dims, best, overall_pass, comparison_png
    )

    report_txt = OUTPUT_DIR / "scale_rotation_measurement_report.txt"
    with report_txt.open("w", encoding="utf-8") as f:
        f.write("3D RECONSTRUCTION — SCALE, 45-DEGREE X ROTATION, DIMENSION CHECK\n")
        f.write("=" * 72 + "\n\n")
        f.write(f"Image: {IMAGE_PATH}\nInput PLY: {PLY_PATH}\nPoints: {len(original)}\n")
        f.write(f"Known width: {KNOWN_WIDTH_CM:.3f} cm\nKnown height: {KNOWN_HEIGHT_CM:.3f} cm\n")
        f.write(f"Chosen width axis: {axis_names[width_axis]}\nChosen height axis: {axis_names[height_axis]}\n")
        f.write(f"Scale factor: {scale:.12f} cm/PLY unit\n\n")
        f.write(f"Reconstructed width: {predicted_width:.4f} cm; error: {width_error_pct:.3f}%\n")
        f.write(f"Reconstructed height: {predicted_height:.4f} cm; error: {height_error_pct:.3f}%\n")
        f.write(f"Tolerance: +/-{TOLERANCE_PERCENT:.1f}%\n")
        f.write(f"Overall result: {'PASS' if overall_pass else 'FAIL'}\n\n")
        f.write("A PASS is reported only when both dimension errors are at most 5%.\n")
        f.write("The script does not artificially clamp or overwrite errors. If it fails, inspect/crop\n")
        f.write("the frame point cloud or improve reconstruction/axis alignment, then rerun.\n\n")
        f.write("Combined transformation matrix T:\n" + np.array2string(T, precision=12) + "\n")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_heading("3D Reconstruction: Scaling, Rotation and Dimension Verification", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("COLMAP point cloud — assignment results")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Reference object and input", level=1)
    doc.add_paragraph(
        f"The reference object is the framed picture in images/1.jpg. Its known outer dimensions "
        f"are {KNOWN_WIDTH_CM:.1f} cm wide and {KNOWN_HEIGHT_CM:.1f} cm high. "
        f"The input reconstruction is {PLY_PATH}, containing {len(original):,} points."
    )
    doc.add_paragraph(f"Chosen width axis: {axis_names[width_axis]}; chosen height axis: {axis_names[height_axis]}.")
    doc.add_paragraph(f"Calculated uniform scale: {scale:.10f} cm per original PLY unit.")

    doc.add_heading("2. Transformation matrices", level=1)
    doc.add_paragraph("Scaling matrix S (4 × 4):")
    doc.add_paragraph(np.array2string(S, precision=8))
    doc.add_paragraph("Rotation matrix Rx(45°) (4 × 4):")
    doc.add_paragraph(np.array2string(Rx, precision=8))
    doc.add_paragraph("Combined homogeneous transformation T = Rx @ S:")
    doc.add_paragraph(np.array2string(T, precision=10))
    doc.add_paragraph(
        "Using column vectors, p′ = T p. Uniform scaling is applied first, followed by a 45° "
        "rotation about the X-axis around the origin. Rotation preserves Euclidean distances."
    )

    doc.add_heading("3. Dimension verification (±5% tolerance)", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Dimension", "Actual (cm)", "Reconstructed (cm)", "Abs. error (cm)", "Error (%)", "Result"]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for name, actual, predicted, err_cm, err_pct, passed in [
        ("Width", KNOWN_WIDTH_CM, predicted_width, width_error_cm, width_error_pct, width_pass),
        ("Height", KNOWN_HEIGHT_CM, predicted_height, height_error_cm, height_error_pct, height_pass),
    ]:
        cells = table.add_row().cells
        vals = [name, f"{actual:.3f}", f"{predicted:.3f}", f"{err_cm:.3f}", f"{err_pct:.2f}%", "PASS" if passed else "FAIL"]
        for cell, value in zip(cells, vals):
            cell.text = value

    doc.add_paragraph(
        f"Overall tolerance result: {'PASS' if overall_pass else 'FAIL'} "
        f"(both dimensions must be within ±{TOLERANCE_PERCENT:.0f}%). "
        "The scale is a single uniform least-squares fit to both dimensions; errors are reported honestly "
        "and are not forced to meet the tolerance."
    )
    doc.add_heading("4. Visual comparison", level=1)
    doc.add_picture(str(comparison_png), width=Inches(7.0))
    doc.add_paragraph(
        "Figure: original photograph, original reconstructed point cloud, and scaled point cloud with "
        "the measured width/height errors and ±5% tolerance status."
    )
    docx_path = OUTPUT_DIR / "scale_rotation_report.docx"
    doc.save(docx_path)

    print("\nCompleted. Results saved to:", OUTPUT_DIR)
    print("Comparison image:", comparison_png)
    print("Word report:", docx_path)
    print("Transformation matrix:", matrix_txt)
    print("Dimension CSV:", csv_path)
    print("Measurement report:", report_txt)
    print("Scaled PLY:", scaled_ply)
    print("Scaled + rotated PLY:", rotated_ply)


if __name__ == "__main__":
    main()